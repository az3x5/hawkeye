from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/home/axmyn/Projects/AI/eagleeye-demo/docs/EagleEye_BlackGlass_API_Integration_Guide.docx"
NAVY = "17365D"
PALE = "EEF4FA"
GRAY = "F4F5F7"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = str(text)
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                set_cell_shading(cells[idx], PALE)
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9.2)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["No Spacing"]
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), GRAY)
    pPr.append(shd)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")
    run.font.size = Pt(8.3)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.8)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12

for style_name, size, before, after in (
    ("Title", 28, 0, 14),
    ("Subtitle", 13, 0, 10),
    ("Heading 1", 17, 16, 7),
    ("Heading 2", 13, 11, 5),
    ("Heading 3", 11, 8, 4),
):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = style_name != "Subtitle"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("EagleEye and BlackGlass API Integration Guide   |   16 September 2026")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(100, 100, 100)

title = doc.add_paragraph(style="Title")
title.add_run("EagleEye and BlackGlass API Integration Guide")
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.add_run("Two way ingestion and cited analysis result delivery")

doc.add_paragraph("Prepared for BlackGlass integration developers and EagleEye operators")
doc.add_paragraph("Version 1.1   |   16 September 2026")
doc.add_paragraph()

p = doc.add_paragraph()
r = p.add_run("Purpose")
r.bold = True
p.add_run(
    "  This guide defines how BlackGlass submits text and media to EagleEye and how "
    "EagleEye returns processed, source-linked analysis to BlackGlass. It includes the "
    "live public ingestion endpoints, required fields, authentication, acknowledgements, "
    "deduplication rules, result delivery, and the remaining activation work."
)

p = doc.add_paragraph()
r = p.add_run("Current conclusion")
r.bold = True
p.add_run(
    "  The public tunnel and the legacy text and media ingestion routes are online. "
    "Automatic result delivery is implemented in the newer local evidence pipeline but is "
    "not yet active on Cyber-AI. BlackGlass must provide an HTTPS result receiver and a "
    "separate receiver token before that delivery worker can be enabled."
)

doc.add_page_break()

doc.add_heading("Integration Overview", level=1)
doc.add_paragraph(
    "BlackGlass remains the collection and presentation system. EagleEye receives authorized "
    "source records, retains provenance, runs bounded analysis, and returns cited results. "
    "BlackGlass must preserve its own object identifiers so every EagleEye result can be "
    "joined back to the correct post, profile, person, vehicle, case, or live segment."
)

add_table(
    doc,
    ["Direction", "Purpose", "Authentication", "Current state"],
    [
        ["BlackGlass to EagleEye", "Submit text and binary evidence", "EagleEye bearer token", "Online for text and media"],
        ["EagleEye to BlackGlass", "Deliver cited analysis snapshots", "Separate BlackGlass receiver token", "Code ready locally but not deployed"],
        ["BlackGlass replay", "Recover missed result events", "EagleEye bearer token", "Part of newer pipeline"],
    ],
    [1.35, 2.15, 1.75, 1.75],
)

doc.add_heading("Security Model", level=1)
add_bullets(
    doc,
    [
        "Use HTTPS only. The EagleEye origin is not exposed directly; Cloudflare Tunnel publishes an exact route allowlist.",
        "Send the EagleEye service token only in the Authorization header. Never place tokens in URLs, form metadata, source URLs, logs, or browser storage.",
        "Use one stable BlackGlass service identity across token rotation so owner-scoped result access remains consistent.",
        "Use a different token for EagleEye result delivery to the BlackGlass receiver.",
        "Treat biometric and restricted records according to their classification and authorization basis.",
        "Store secrets in server-side secret management. Do not commit them to source control or include them in support documents.",
    ],
)

doc.add_heading("Public EagleEye Base URL", level=1)
add_code(doc, "https://api.staticdev.xyz")
doc.add_paragraph("All public requests require the following header:")
add_code(doc, "Authorization: Bearer <EAGLEEYE_SERVICE_TOKEN>")

doc.add_heading("Live Public Endpoints", level=1)
add_table(
    doc,
    ["Method", "Path", "Purpose", "Required scope"],
    [
        ["GET", "/api/v1/integrations/blackglass/capabilities", "Read supported ingestion capabilities", "media:write"],
        ["POST", "/api/v1/integrations/blackglass/text", "Submit a text record", "language"],
        ["POST", "/api/v1/integrations/blackglass/media", "Submit one binary object", "media:write"],
        ["POST", "/api/v1/live/frames/analyze", "Analyze one finite live frame", "Configured identification access"],
    ],
    [0.65, 3.25, 2.05, 1.05],
)

doc.add_heading("Text Ingestion", level=1)
doc.add_paragraph(
    "Use the text endpoint for posts, captions, comments, messages, articles, and other UTF-8 "
    "content. EagleEye accepts English, Dhivehi in Thaana, Latin-script Dhivehi, and mixed text. "
    "The source ID must be stable and unique within its source type."
)
add_code(
    doc,
    "POST /api/v1/integrations/blackglass/text\n"
    "Authorization: Bearer <EAGLEEYE_SERVICE_TOKEN>\n"
    "Content-Type: application/json",
)

doc.add_heading("Text Request Fields", level=2)
add_table(
    doc,
    ["Field", "Required", "Type and limit", "Meaning"],
    [
        ["schema_version", "No", "String fixed at 1.1", "Contract version"],
        ["source_id", "Yes", "String 1 to 256", "Stable BlackGlass record ID"],
        ["source_type", "Yes", "String 1 to 64", "post, profile, media, or another record type"],
        ["title", "Yes", "String 1 to 256", "Short display title"],
        ["text", "Yes", "String 1 to 100000", "Original text without replacement"],
        ["language_hint", "No", "String up to 16", "For example en, dv, latin-dv, or mixed"],
        ["attributes", "No", "JSON object", "Profile, platform, case, report, or collector metadata"],
        ["requested_analyses", "No", "Unique string array", "Requested processing capabilities"],
    ],
    [1.35, 0.65, 1.55, 3.45],
)

doc.add_heading("Text Request Example", level=2)
add_code(
    doc,
    '''{
  "schema_version": "1.1",
  "source_id": "POST-1001",
  "source_type": "post",
  "attributes": {
    "source_system": "blackglass-prod",
    "source_url": "https://blackglass.live/posts/POST-1001",
    "collected_at": "2026-09-15T10:00:00Z",
    "published_at": "2026-09-15T09:30:00Z",
    "collector_version": "1.0",
    "profile_id": "PROFILE-456",
    "platform": "facebook",
    "report_request_id": "REPORT-123"
  },
  "title": "Post title",
  "text": "Original English, Thaana, Latin Dhivehi or mixed text",
  "language_hint": "mixed",
  "requested_analyses": [
    "semantic_embedding", "translation", "transliteration"
  ]
}''',
)

doc.add_heading("Media Ingestion", level=1)
doc.add_paragraph(
    "Use multipart upload for images, audio, video, PDF documents, and post attachments. "
    "BlackGlass sends the actual bytes; a source URL is provenance and EagleEye does not fetch it."
)
add_code(
    doc,
    "POST /api/v1/integrations/blackglass/media\n"
    "Authorization: Bearer <EAGLEEYE_SERVICE_TOKEN>\n"
    "Content-Type: multipart/form-data",
)

doc.add_heading("Media Form Fields", level=2)
add_table(
    doc,
    ["Field", "Required", "Accepted value", "Meaning"],
    [
        ["file", "Yes", "Binary upload", "Original evidence bytes"],
        ["source_id", "Yes", "String 1 to 256", "Stable BlackGlass object ID"],
        ["source_type", "Yes", "String 1 to 64", "image, video, audio, document, attachment"],
        ["attributes", "No", "JSON object", "Optional source metadata not used for exact lookup"],
        ["source_system", "No", "Default blackglass-prod", "Source namespace"],
        ["classification", "No", "public, internal, restricted, biometric", "Sensitivity class"],
        ["source_url", "No", "URL up to 2048", "Original record location for provenance"],
        ["collected_at", "No", "ISO 8601 datetime", "Collection time with timezone"],
        ["published_at", "No", "ISO 8601 datetime", "Publication time with timezone"],
        ["collector_version", "No", "String up to 128", "Collector release identifier"],
        ["requested_analyses", "No", "CSV or JSON array", "Requested processing capabilities"],
    ],
    [1.55, 0.65, 2.1, 2.9],
)

doc.add_heading("Media Request Example", level=2)
add_code(
    doc,
    '''curl -X POST \\
  https://api.staticdev.xyz/api/v1/integrations/blackglass/media \\
  -H "Authorization: Bearer $EAGLEEYE_SERVICE_TOKEN" \\
  -F "source_id=MEDIA-1001" \\
  -F "source_type=image" \\
  -F "source_system=blackglass-prod" \\
  -F "classification=internal" \\
  -F "source_url=https://blackglass.live/media/MEDIA-1001" \\
  -F "collected_at=2026-09-15T10:00:00Z" \\
  -F "requested_analyses=object_detection,landmark_recognition,ocr" \\
  -F "file=@photo.jpg"''',
)

doc.add_heading("Supported Analysis Names", level=1)
add_table(
    doc,
    ["Category", "Capabilities"],
    [
        ["Faces", "face_detection, face_identification"],
        ["Scenes", "object_detection, landmark_recognition, video_tracking"],
        ["Vehicles", "vehicle_detection, vehicle_identification"],
        ["Language", "ocr, transcription, translation, transliteration, semantic_embedding"],
    ],
    [1.25, 5.75],
)
doc.add_paragraph(
    "A capability response may report queued, available_on_demand, not_connected, or "
    "not_applicable. Installed model files do not by themselves mean that an automatic worker "
    "is connected. BlackGlass must read the returned state rather than assume completion."
)

doc.add_page_break()
doc.add_heading("Ingestion Acknowledgement", level=1)
doc.add_paragraph(
    "EagleEye returns HTTP 202 after it has accepted or recognized the submitted object. "
    "BlackGlass must store subject.id beside the original BlackGlass object ID. A repeated "
    "delivery may return already_exists and must be treated as successful deduplication."
)
add_code(
    doc,
    '''{
  "schema_version": "1.1",
  "source_id": "MEDIA-1001",
  "source_type": "image",
  "attributes": {
    "source_system": "blackglass-prod"
  },
  "subject": {
    "type": "media",
    "id": "8be36c1e-4d58-4d80-81f9-3adb2be86262"
  },
  "status": "accepted",
  "created": true,
  "content_type": "image/jpeg",
  "analysis_routes": [
    {
      "capability": "object_detection",
      "state": "not_connected",
      "endpoint": null,
      "detail": "Automatic worker is not connected."
    }
  ]
}''',
)

doc.add_heading("Status and Error Handling", level=1)
add_table(
    doc,
    ["HTTP status", "Meaning", "BlackGlass action"],
    [
        ["202", "Accepted or deduplicated", "Persist the acknowledgement and object mapping"],
        ["401", "Token missing, invalid, or expired", "Stop and refresh or replace the credential"],
        ["403", "Token lacks the required scope", "Request the narrow required service scope"],
        ["413", "Upload exceeds the configured limit", "Split supported content or use the approved bulk path"],
        ["422", "Invalid fields, type, or analysis name", "Correct the record; do not retry unchanged"],
        ["429", "Owner queue is at capacity", "Retry with backoff and jitter"],
        ["503", "Required storage or processing service unavailable", "Retry with bounded exponential backoff"],
    ],
    [0.85, 2.25, 3.9],
)

doc.add_heading("Sender Reliability Rules", level=1)
add_bullets(
    doc,
    [
        "Persist each successful acknowledgement before advancing the BlackGlass export cursor.",
        "Reuse stable source_id and source_type values on retries.",
        "Treat accepted and already_exists as successful outcomes.",
        "Apply bounded exponential backoff with jitter for 429, 502, 503, timeouts, and connection failures.",
        "Do not retry unchanged 401, 403, 413, or 422 requests.",
        "Keep the original text and timestamps. Do not replace the submitted source with a translation or normalized form.",
    ],
)

doc.add_heading("EagleEye Result Delivery", level=1)
doc.add_paragraph(
    "BlackGlass must expose one HTTPS endpoint that receives complete analysis snapshots. "
    "The recommended route is shown below; BlackGlass may choose another fixed HTTPS URL and "
    "provide it to the EagleEye operator."
)
add_code(doc, "POST https://blackglass.live/api/v1/integrations/eagleeye/results")

doc.add_heading("Result Request Headers", level=2)
add_code(
    doc,
    "Authorization: Bearer <BLACKGLASS_RECEIVER_TOKEN>\n"
    "Idempotency-Key: <event_id>\n"
    "Content-Type: application/json",
)

doc.add_heading("Result Event Fields", level=2)
add_table(
    doc,
    ["Field", "Purpose"],
    [
        ["event_id", "Stable delivery ID and idempotency key"],
        ["event_type", "analysis.updated"],
        ["analysis_id", "Stable EagleEye analysis ID"],
        ["report_request_id", "BlackGlass report correlation ID"],
        ["subject", "Profile, person, vehicle, case, or other report subject"],
        ["revision", "Monotonic revision for the analysis"],
        ["status", "queued, running, retry, completed, partial, or failed"],
        ["source_id", "Original BlackGlass record ID, indexed for exact lookup"],
        ["source_type", "Original BlackGlass record type, indexed with source_id"],
        ["attributes", "Optional source URL, system, collection time, and other metadata"],
        ["stream", "Optional finite live segment identity and timing"],
        ["source_sha256", "Checksum of the submitted original"],
        ["evidence", "Extracted text and observations with exact source locators"],
        ["findings", "Unreviewed findings that cite evidence IDs"],
        ["contradictions", "Conflicts found in the analyzed evidence"],
        ["warnings", "Unsupported, partial, truncated, or uncertain processing notes"],
        ["report", "BlackGlass compatible schema 2 report derived from cited evidence"],
    ],
    [1.65, 5.35],
)

doc.add_heading("Result Event Example", level=2)
add_code(
    doc,
    '''{
  "schema_version": "1.1",
  "event_id": "0bb91572-c924-43fa-a935-b2222688435c",
  "event_type": "analysis.updated",
  "emitted_at": "2026-09-15T10:05:00Z",
  "analysis_id": "1a151bcb-d139-4dba-b36a-99bf268794ef",
  "report_request_id": "REPORT-123",
  "subject": {
    "subject_type": "social_profile",
    "subject_id": "PROFILE-456",
    "display_label": "Profile under review"
  },
  "revision": 1,
  "status": "completed",
  "source_id": "POST-1001",
  "source_type": "post",
  "attributes": {
    "source_system": "blackglass-prod",
    "source_url": "https://blackglass.live/posts/POST-1001"
  },
  "source_sha256": "<64 lowercase hexadecimal characters>",
  "evidence": [
    {
      "evidence_id": "571539dd-6626-464b-803d-c659b23aad69",
      "kind": "ocr",
      "original_text": "Original extracted text",
      "normalized_text": "Normalized text",
      "locator": {"page": 1, "precision": "source"},
      "provenance": {"model": "model-name", "version": "model-version"},
      "translations": [],
      "review_status": "unreviewed"
    }
  ],
  "findings": [],
  "contradictions": [],
  "warnings": [],
  "report": {"schemaVersion": "2.0", "document": {}}
}''',
)

doc.add_page_break()
doc.add_heading("BlackGlass Result Acknowledgement", level=1)
doc.add_paragraph(
    "BlackGlass must durably store the event before returning success. The response event_id "
    "must exactly match the request. EagleEye does not follow redirects and treats a wrong or "
    "missing acknowledgement as a failed delivery."
)
doc.add_heading("First Delivery", level=2)
add_code(doc, '{"event_id":"<same-event-uuid>","status":"accepted"}')
doc.add_heading("Replay Already Stored", level=2)
add_code(doc, '{"event_id":"<same-event-uuid>","status":"duplicate"}')
doc.add_paragraph(
    "EagleEye retries the same event ID after network or acknowledgement failures. BlackGlass "
    "must process each event ID once and return duplicate for an already stored event."
)

doc.add_heading("Source and Citation Rules", level=1)
add_bullets(
    doc,
    [
        "Join results using source_id and source_type. Do not scan attributes or rely only on labels.",
        "Display findings as unreviewed until an authorized reviewer confirms them.",
        "Every finding must cite evidence IDs contained in the same committed analysis revision.",
        "Keep original text separate from normalized text, translations, and transliterations.",
        "Use evidence locators to open the relevant page, character range, audio interval, video interval, or sampled frame.",
        "A citation proves traceability to an extraction. It does not prove that the extraction or model interpretation is factually correct.",
        "Do not automatically publish alerts or identity claims from unreviewed findings.",
    ],
)

doc.add_heading("Live Data", level=1)
doc.add_paragraph(
    "The live API analyzes one submitted frame at a time. Continuous sources must be segmented "
    "by the capture system. Each segment or frame needs a stable stream ID, segment ID, sequence, "
    "and timezone-aware start time when the newer evidence workflow is enabled. This supports "
    "incremental processing; it is not a direct RTSP or WebRTC capture server."
)
add_code(
    doc,
    '''{
  "stream_id": "CAMERA-01",
  "segment_id": "CAMERA-01-000042",
  "sequence": 42,
  "started_at": "2026-09-15T10:04:30Z"
}''',
)

doc.add_page_break()
doc.add_heading("Activation Status", level=1)
add_table(
    doc,
    ["Component", "Status", "Action before production"],
    [
        ["Cloudflare tunnel", "Online", "Monitor connector health and certificate routing"],
        ["Text ingestion", "Online", "Issue a scoped BlackGlass service token"],
        ["Media ingestion", "Online", "Issue a scoped BlackGlass service token and test size limits"],
        ["Live frame route", "Online", "Test with representative camera frames and authorization"],
        ["Unified evidence ingestion", "Ready locally", "Deploy the newer evidence API and database migration"],
        ["Evidence analysis worker", "Ready locally", "Deploy and start the explicit evidence worker profile"],
        ["Result replay API", "Ready locally", "Deploy and expose only the exact owner-scoped routes"],
        ["Automatic webhook delivery", "Not active", "Obtain BlackGlass receiver URL and token, then enable delivery"],
        ["Cloudflare Access", "Not active", "Optional second machine-authentication layer requires Cloudflare Access write permission"],
        ["Language quality release", "Not signed off", "Validate reviewed Dhivehi, English, Latin Dhivehi, and mixed samples"],
    ],
    [2.05, 1.2, 3.75],
)

doc.add_heading("Information BlackGlass Must Provide", level=1)
add_bullets(
    doc,
    [
        "The exact HTTPS result receiver URL.",
        "A receiver bearer token stored as a server secret.",
        "Confirmation that the receiver validates Idempotency-Key and returns the required JSON acknowledgement.",
        "The stable BlackGlass service identity that will own submitted analyses.",
        "Expected source object types, maximum file sizes, volume, and delivery concurrency.",
        "The mapping from report_request_id and subject IDs to BlackGlass reports and profiles.",
        "Representative authorized test records for Dhivehi, English, mixed text, images, audio, and video.",
    ],
)

doc.add_heading("Production Acceptance Checklist", level=1)
add_bullets(
    doc,
    [
        "Confirm the service token has language, media:write, and media:read only as required.",
        "Test text, image, audio, video, and document ingestion with real BlackGlass identifiers.",
        "Verify accepted and already_exists handling and durable sender cursor updates.",
        "Deploy the newer evidence migration, API, analysis worker, and result replay route.",
        "Complete an end-to-end result webhook test, including lost acknowledgement and duplicate replay.",
        "Verify that BlackGlass links every finding to the correct source object and evidence locator.",
        "Measure queue capacity, processing latency, storage growth, and retry behavior at the intended load.",
        "Evaluate false identity matches and representative language accuracy before enabling unattended alerts.",
        "Back up PostgreSQL and retained originals together and test restoration.",
    ],
)

doc.add_heading("Contact Values for Configuration", level=1)
add_table(
    doc,
    ["Setting", "Value"],
    [
        ["EagleEye public base URL", "https://api.staticdev.xyz"],
        ["EagleEye service token", "Exchange securely outside this document"],
        ["BlackGlass result receiver URL", "To be provided by BlackGlass"],
        ["BlackGlass receiver token", "Exchange securely outside this document"],
        ["Contract version", "1.1"],
    ],
    [2.35, 4.65],
)

doc.core_properties.title = "EagleEye and BlackGlass API Integration Guide"
doc.core_properties.subject = "Two way ingestion and cited analysis result delivery"
doc.core_properties.author = "EagleEye Project"
doc.core_properties.keywords = "EagleEye, BlackGlass, API, ingestion, evidence, webhook"
doc.save(OUT)
print(OUT)
